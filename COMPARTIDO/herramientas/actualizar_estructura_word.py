# Genera ESTRUCTURA_ORGANIZATIVA_LOGISTICA_INVENTARIOS_ERP_ACTUALIZADO.docx a partir del Word entregado por el usuario,
# corrigiéndolo con las plantillas Excel (que MANDAN: son lo más actual) y con las decisiones cerradas del proyecto:
#   J1 no existe el tipo de movimiento Ajuste (regularizar = ingreso o salida con motivo)
#   J2 producto fallado = salida + ingreso · J3 tercerización en la OF · T2/T7 transferencia en dos pasos (aprobar compromete origen)
#   2026-09-16 se crea el almacén de producto en proceso (SB-ZARATE-PP / CN-ZARATE-PP).
# Uso (desde PROTOTIPOS):  python COMPARTIDO/herramientas/actualizar_estructura_word.py
import copy, os, sys, openpyxl, docx
from docx.text.paragraph import Paragraph

sys.stdout.reconfigure(encoding='utf-8')
BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
PL = os.path.join(BASE, 'docs', 'INFO', 'PLANTILLAS ENTREGADAS POR EL USUARIO')
ORIGEN = os.path.join(PL, 'ESTRUCTURA_ORGANIZATIVA_LOGISTICA_INVENTARIOS_ERP.docx')
DESTINO = os.path.join(PL, 'ESTRUCTURA_ORGANIZATIVA_LOGISTICA_INVENTARIOS_ERP_ACTUALIZADO.docx')


def xl(archivo, hoja):
    ws = openpyxl.load_workbook(os.path.join(PL, archivo), data_only=True)[hoja]
    return [['' if c is None else str(c).strip() for c in r] for r in ws.iter_rows(values_only=True)]


def bloque(archivo, titulo):
    out, dentro = [], False
    for r in xl(archivo, 'Configuraciones'):
        v = [c for c in r if c]
        if not v:
            continue
        if v[0][:2].rstrip('.').isdigit() and len(v) == 1:
            dentro = titulo.upper() in v[0].upper()
            continue
        if dentro:
            out.append(v)
    return out[1:]  # sin cabecera


def celdas(tr):
    return list(dict.fromkeys(tc for tc in tr.iter(docx.oxml.ns.qn('w:tc'))))


def poner_texto(tc, texto, doc):
    ps = tc.findall(docx.oxml.ns.qn('w:p'))
    for p in ps[1:]:
        tc.remove(p)
    par = Paragraph(ps[0], doc)
    runs = par.runs
    if runs:
        runs[0].text = texto
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        par.add_run(texto)


def llenar(tabla, filas, doc):
    """Deja la cabecera y reemplaza las filas de datos (usa la primera fila de datos como molde)."""
    trs = tabla._tbl.tr_lst
    molde = copy.deepcopy(trs[1])
    for tr in trs[1:]:
        tabla._tbl.remove(tr)
    for f in filas:
        tr = copy.deepcopy(molde)
        tcs = celdas(tr)
        for i, tc in enumerate(tcs):
            poner_texto(tc, f[i] if i < len(f) else '', doc)
        tabla._tbl.append(tr)


def parrafo_despues(elemento, texto, doc, negrita=False):
    p = copy.deepcopy(doc.paragraphs[0]._p)
    for r in p.findall(docx.oxml.ns.qn('w:r')):
        p.remove(r)
    elemento.addnext(p)
    par = Paragraph(p, doc)
    run = par.add_run(texto)
    run.bold = negrita
    return p


doc = docx.Document(ORIGEN)
T = doc.tables

# ---------- Excel ----------
alm = [r for r in xl('PLANTILLA_Almacenes_GI.xlsx', 'Almacenes')[1:] if r[1]]
empresas = bloque('PLANTILLA_Almacenes_GI.xlsx', '1. EMPRESAS')
sedes = bloque('PLANTILLA_Almacenes_GI.xlsx', '2. SEDES')
grupos_prov = bloque('PLANTILLA_Proveedores.xlsx', '3. GRUPO DE PROVEEDORES')
ALM_PP = {'SB': ['SB', 'SB-ZARATE-PP', 'Almacén Zárate Producto en Proceso', 'Común', 'Zárate', 'Físico', 'Productos en Proceso'],
          'CN': ['CN', 'CN-ZARATE-PP', 'Almacén Zárate Producto en Proceso', 'Común', 'Zárate', 'Físico', 'Productos en Proceso']}

# 0 Centros
llenar(T[0], [[e[0], e[1], e[2], e[3]] for e in empresas], doc)
# 1 Sedes (la columna Compartida no está en el Excel: se conserva la del Word; tiendas = No)
compartida = {'G': 'Sí', 'Z': 'Sí', 'VIRT': 'Sí'}
llenar(T[1], [[s[0], s[1] + (' (' + s[2] + ')' if len(s) > 2 and s[2] and not s[2].startswith('(') else ''), compartida.get(s[0], 'No'), s[3] if len(s) > 3 else ''] for s in sedes], doc)
# 2 Almacenes comunes (SB y CN en paralelo) + almacén de producto en proceso
comunes_sb = [r for r in alm if r[0] == 'SB' and r[3] == 'Común']
filas = []
for r in comunes_sb:
    cn = next((x for x in alm if x[0] == 'CN' and x[1] == r[1].replace('SB-', 'CN-')), None)
    filas.append([r[1], cn[1] if cn else '', r[2], r[4], r[6]])
    if r[1] == 'SB-ZARATE-MP':
        filas.append([ALM_PP['SB'][1], ALM_PP['CN'][1], ALM_PP['SB'][2], 'Zárate', 'Productos en Proceso (nuevo 2026-09-16)'])
llenar(T[2], filas, doc)
# 3 Transición
llenar(T[3], [[r[1], r[2], r[4]] for r in alm if r[3] == 'Transición'], doc)
# 4 y 5 Tiendas (galería según Excel)
llenar(T[4], [[r[1], r[2], r[4]] for r in alm if r[0] == 'SB' and r[3].startswith('Tienda')], doc)
llenar(T[5], [[r[1], r[2], r[4]] for r in alm if r[0] == 'CN' and r[3].startswith('Tienda')], doc)
# 8 Grupos de proveedores (Excel)
llenar(T[8], [[g[0], g[1] + (' · ' + g[2] if len(g) > 2 else '')] for g in grupos_prov], doc)
# 9 Grupos de movimiento: sin Ajustes (J1)
llenar(T[9], [r for r in [[c.text for c in row.cells[:3]] for row in T[9].rows][1:] if r[0] != 'AJU'], doc)
# 10 Ingresos
ing = [[c.text for c in row.cells[:3]] for row in T[10].rows][1:]
ing += [['ING-REGULARIZ', 'Ingreso por regularización de inventario (sobrante)', 'Conteo > sistema -> Almacén'],
        ['ING-OBSERV', 'Ingreso con observación (recepción no conforme)', 'Proveedor -> Almacén (deja observación trazada)'],
        ['ING-FALLADO', 'Ingreso del artículo fallado (producto fallado)', 'Artículo original -> Artículo FALLADO, mismo almacén y costo']]
llenar(T[10], ing, doc)
# 11 Salidas
sal = [[c.text for c in row.cells[:3]] for row in T[11].rows][1:]
sal += [['SAL-REGULARIZ', 'Salida por regularización de inventario (faltante)', 'Almacén -> Conteo < sistema'],
        ['SAL-FALLADO', 'Salida por producto fallado', 'Artículo original -> Artículo FALLADO']]
llenar(T[11], sal, doc)
# 12 Transferencias: TRF-FABRIC según J3
trf = [[c.text for c in row.cells[:3]] for row in T[12].rows][1:]
for r in trf:
    if r[0] == 'TRF-FABRIC':
        r[2] = 'Almacén propio -> Almacén en Transición (proveedor del servicio); el retorno es el recibo de producción de la OF'
llenar(T[12], trf, doc)

# ---------- quitar Ajustes (J1) y anotar las decisiones ----------
tabla_aju = T[13]._tbl
titulo_aju = tabla_aju.getprevious()
while titulo_aju is not None and 'Ajustes' not in ''.join(titulo_aju.itertext()):
    titulo_aju = titulo_aju.getprevious()
ancla = titulo_aju.getprevious()
tabla_aju.getparent().remove(tabla_aju)
titulo_aju.getparent().remove(titulo_aju)
notas = [
    ('Decisiones cerradas que aplican a los movimientos', True),
    ('J1 · No existe el grupo ni el tipo de movimiento Ajuste. Regularizar inventario es un Ingreso (ING-REGULARIZ) o una Salida (SAL-REGULARIZ) con motivo «Regularización de inventario (sobrante / faltante)» y observación, sin visto bueno.', False),
    ('J2 · Producto fallado = Salida del artículo (SAL-FALLADO) + Ingreso del artículo «… FALLADO» (ING-FALLADO) al mismo costo.', False),
    ('T2 / T7 · Transferencia en dos pasos (GI-11): al aprobarse compromete el stock en el almacén de origen y suma «Pedido» en el destino; al confirmar la recepción sale del origen y entra al destino. Se puede recibir por partes y cancelar lo pendiente.', False),
    ('J3 · La tercerización vive en la Orden de Fabricación: el envío al proveedor es una transferencia TRF-FABRIC al almacén en Transición; el retorno es el recibo de producción; el servicio se compra con una OC de servicio normal.', False),
]
p = ancla
for texto, neg in notas:
    p = parrafo_despues(p, texto, doc, neg)

# ---------- encabezado de la versión ----------
primero = doc.paragraphs[0]._p
nota = parrafo_despues(primero, 'Versión actualizada 2026-09-16. Mandan las plantillas Excel (almacenes, sedes, empresas y grupos de proveedores); se agregó el almacén de producto en proceso y se aplicaron las decisiones cerradas J1, J2, J3, T2 y T7. El Word original queda como referencia.', doc)
primero.addprevious(nota)

doc.save(DESTINO)
print('Generado', DESTINO)
